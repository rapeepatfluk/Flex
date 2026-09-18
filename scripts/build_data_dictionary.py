from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCHEMA = ROOT / "database" / "schema_latest.sql"
OUTPUT = ROOT / "docs" / "FLEXJOB_Data_Dictionary_Thesis_Format.docx"

TABLE_INFO = {
    "applications": ("APPLICATIONS", "ใบสมัครงาน", "เก็บข้อมูลการสมัครงาน สถานะการพิจารณา และคะแนนหลังจบงาน"),
    "auth_tokens": ("AUTH TOKENS", "โทเคนยืนยันตัวตน", "เก็บโทเคนสำหรับยืนยันอีเมลและตั้งรหัสผ่านใหม่"),
    "email_log": ("EMAIL LOG", "คิวและประวัติอีเมล", "เก็บอีเมลที่รอส่ง กำลังส่ง ส่งสำเร็จ หรือส่งไม่สำเร็จ"),
    "employer_documents": ("EMPLOYER DOCUMENTS", "เอกสารผู้ว่าจ้าง", "เก็บเอกสารยืนยันตัวตนหรือบริษัทของผู้ว่าจ้าง"),
    "employer_profiles": ("EMPLOYER PROFILES", "โปรไฟล์ผู้ว่าจ้าง", "เก็บข้อมูลบริษัทของบัญชีผู้ว่าจ้าง"),
    "job_categories": ("JOB CATEGORIES", "ประเภทงาน", "เก็บประเภทงานหลัก เช่น พาร์ทไทม์ อีเวนต์ และฟรีแลนซ์"),
    "job_images": ("JOB IMAGES", "รูปประกาศงาน", "เก็บรูปภาพประกอบของประกาศงาน"),
    "job_invitations": ("JOB INVITATIONS", "คำเชิญสมัครงาน", "เก็บคำเชิญที่ผู้ว่าจ้างส่งให้ผู้หางาน"),
    "job_promotions": ("JOB PROMOTIONS", "รายการโปรโมตประกาศ", "เก็บการสั่งซื้อ การชำระเงิน และสถานะการโปรโมตประกาศ"),
    "job_skills": ("JOB SKILLS", "ทักษะของประกาศงาน", "เชื่อมประกาศงานกับทักษะที่ต้องการ"),
    "job_worker_matches": ("JOB WORKER MATCHES", "ผลการจับคู่งานและผู้หางาน", "เก็บผล Matching แบบ cache เพื่อให้ค้นหาและแสดงผลเร็วขึ้น"),
    "jobs": ("JOBS", "ประกาศงาน", "เก็บรายละเอียดประกาศงานทั้งหมด"),
    "notifications": ("NOTIFICATIONS", "การแจ้งเตือน", "เก็บการแจ้งเตือนภายในระบบของผู้ใช้"),
    "promotion_packages": ("PROMOTION PACKAGES", "แพ็กเกจโปรโมต", "เก็บแพ็กเกจและราคาในการดันประกาศงาน"),
    "schema_migrations": ("SCHEMA MIGRATIONS", "ประวัติ Migration", "เก็บรายการ migration ที่ฐานข้อมูลใช้งานแล้ว"),
    "skill_categories": ("SKILL CATEGORIES", "หมวดทักษะ", "เก็บหมวดหมู่ของทักษะมาตรฐาน"),
    "skill_consolidation_map": ("SKILL CONSOLIDATION MAP", "แผนที่รวมทักษะ", "เก็บการโยงทักษะเดิมไปยังทักษะมาตรฐานแบบกว้าง"),
    "skills": ("SKILLS", "ทักษะ", "เก็บรายการทักษะมาตรฐานและทักษะที่ผู้ใช้เพิ่มเอง"),
    "users": ("USERS", "ผู้ใช้งาน", "เก็บข้อมูลบัญชีผู้ดูแล ผู้ว่าจ้าง และผู้หางาน"),
    "work_interests": ("WORK INTERESTS", "ความสนใจด้านงาน", "เก็บหัวข้องานที่ผู้หางานสนใจ"),
    "worker_job_preferences": ("WORKER JOB PREFERENCES", "ประเภทงานที่ผู้หางานสนใจ", "เชื่อมผู้หางานกับประเภทงานที่ต้องการ"),
    "worker_profiles": ("WORKER PROFILES", "โปรไฟล์ผู้หางาน", "เก็บข้อมูลโปรไฟล์ เอกสาร และความพร้อมทำงานของผู้หางาน"),
    "worker_skills": ("WORKER SKILLS", "ทักษะผู้หางาน", "เชื่อมผู้หางานกับทักษะที่มี"),
    "worker_work_interests": ("WORKER WORK INTERESTS", "ความสนใจงานของผู้หางาน", "เชื่อมผู้หางานกับหัวข้องานที่สนใจ"),
}

FIELD_DESCRIPTIONS = {
    "application_id": "รหัสใบสมัครงาน",
    "application_status": "สถานะใบสมัคร เช่น รอพิจารณา สัมภาษณ์ ผ่านงาน หรือไม่ผ่าน",
    "auth_token_id": "รหัสโทเคนยืนยันตัวตน",
    "available_at": "วันและเวลาที่อีเมลพร้อมให้ตัวประมวลผลส่ง",
    "available_from": "วันที่ผู้หางานพร้อมเริ่มงาน",
    "broad_skill_id": "รหัสทักษะมาตรฐานแบบกว้างที่ใช้แทนทักษะเดิม",
    "calculated_at": "วันและเวลาที่คำนวณผล Matching ล่าสุด",
    "category_name": "ชื่อหมวดหมู่",
    "category_slug": "รหัสชื่อหมวดหมู่สำหรับใช้ในระบบ",
    "checksum": "ค่า Hash สำหรับตรวจว่าไฟล์ migration ไม่ถูกแก้ไข",
    "company_address": "ที่อยู่บริษัท",
    "company_description": "รายละเอียดหรือคำอธิบายบริษัท",
    "company_logo_path": "ตำแหน่งไฟล์โลโก้บริษัท",
    "company_name": "ชื่อบริษัท",
    "cover_note": "ข้อความแนะนำตัวที่แนบมากับใบสมัคร",
    "created_at": "วันและเวลาที่สร้างข้อมูล",
    "data_strength": "ความครบถ้วนของข้อมูลที่ใช้คำนวณ Matching",
    "display_order": "ลำดับการแสดงผล",
    "display_priority": "ลำดับความสำคัญในการแสดงแพ็กเกจ",
    "document_file_path": "ตำแหน่งไฟล์เอกสารที่อัปโหลด",
    "document_status": "สถานะการตรวจเอกสาร",
    "duration_days": "จำนวนวันที่สิทธิ์โปรโมตมีผล",
    "email": "อีเมลผู้ใช้งาน",
    "email_verified_at": "วันและเวลาที่ยืนยันอีเมลแล้ว",
    "employer_document_id": "รหัสเอกสารผู้ว่าจ้าง",
    "employer_profile_id": "รหัสโปรไฟล์ผู้ว่าจ้าง",
    "employer_user_id": "รหัสผู้ใช้งานของผู้ว่าจ้าง",
    "ends_at": "วันและเวลาที่สิทธิ์โปรโมตสิ้นสุด",
    "error_msg": "รายละเอียดข้อผิดพลาดในการส่งอีเมล",
    "expires_at": "วันและเวลาหมดอายุของโทเคน",
    "first_name": "ชื่อจริง",
    "html_body": "เนื้อหาอีเมลรูปแบบ HTML ที่รอส่ง",
    "image_file_path": "ตำแหน่งไฟล์รูปภาพ",
    "importance": "ระดับความสำคัญของทักษะ เช่น จำเป็นหรือเสริม",
    "interest_name": "ชื่อหัวข้องานที่สนใจ",
    "interest_slug": "รหัสชื่อหัวข้องานสำหรับใช้ในระบบ",
    "is_active": "ระบุว่ายังเปิดใช้งานข้อมูลนี้หรือไม่",
    "is_custom": "ระบุว่าเป็นทักษะที่ผู้ใช้เพิ่มเองหรือไม่",
    "is_read": "ระบุว่าอ่านการแจ้งเตือนแล้วหรือไม่",
    "job_category_id": "รหัสประเภทงาน",
    "job_description": "รายละเอียดของงาน",
    "job_id": "รหัสประกาศงาน",
    "job_image_id": "รหัสรูปประกาศงาน",
    "job_invitation_id": "รหัสคำเชิญสมัครงาน",
    "job_status": "สถานะประกาศงาน เช่น เปิดรับ ซ่อน หรือปิดรับ",
    "job_title": "ชื่อตำแหน่งหรือชื่อประกาศงาน",
    "legacy_skill_id": "รหัสทักษะเดิมก่อนรวมเป็นทักษะมาตรฐาน",
    "locked_at": "วันและเวลาที่คิวอีเมลถูกล็อกเพื่อส่ง",
    "last_name": "นามสกุล",
    "match_reasons_json": "เหตุผลที่ทำให้คะแนน Matching สูงในรูปแบบ JSON",
    "match_score": "คะแนนความเหมาะสมระหว่างผู้หางานกับประกาศงาน",
    "matching_survey_completed_at": "วันและเวลาที่ตอบแบบสำรวจ Matching แล้ว",
    "matching_survey_required_at": "วันและเวลาที่ระบบกำหนดให้พบแบบสำรวจ Matching",
    "missing_required_json": "ทักษะจำเป็นที่ยังไม่ตรงกันในรูปแบบ JSON",
    "notification_id": "รหัสการแจ้งเตือน",
    "notification_message": "ข้อความรายละเอียดการแจ้งเตือน",
    "notification_title": "หัวข้อการแจ้งเตือน",
    "notification_url": "ลิงก์ปลายทางเมื่อกดการแจ้งเตือน",
    "open_positions": "จำนวนตำแหน่งที่เปิดรับ",
    "package_code": "รหัสแพ็กเกจโปรโมต",
    "package_description": "รายละเอียดแพ็กเกจโปรโมต",
    "package_id": "รหัสแพ็กเกจโปรโมต",
    "package_name": "ชื่อแพ็กเกจโปรโมต",
    "package_name_snapshot": "ชื่อแพ็กเกจที่บันทึกไว้ ณ เวลาสั่งซื้อ",
    "password_hash": "รหัสผ่านที่เข้ารหัสแล้ว",
    "pay_amount": "จำนวนค่าจ้าง",
    "pay_unit": "หน่วยค่าจ้าง เช่น ชั่วโมง วัน หรือโครงการ",
    "payment_reference": "เลขอ้างอิงการชำระเงิน",
    "payment_slip_path": "ตำแหน่งไฟล์หลักฐานการชำระเงิน",
    "payment_submitted_at": "วันและเวลาที่ส่งหลักฐานการชำระเงิน",
    "phone": "หมายเลขโทรศัพท์",
    "portfolio_file_path": "ตำแหน่งไฟล์ Portfolio",
    "portfolio_url": "ลิงก์ Portfolio ภายนอก",
    "preferred_skills_json": "ทักษะเสริมของงานในรูปแบบ JSON",
    "preferred_work_mode": "รูปแบบงานที่ผู้หางานต้องการ",
    "price": "ราคาแพ็กเกจโปรโมต",
    "professional_headline": "ข้อความสรุปสายงานหรือความเชี่ยวชาญ",
    "profile_image_path": "ตำแหน่งไฟล์รูปโปรไฟล์",
    "profile_visibility": "สิทธิ์การแสดงโปรไฟล์ต่อผู้ว่าจ้าง",
    "promotion_id": "รหัสรายการโปรโมต",
    "promotion_status": "สถานะการโปรโมตและการตรวจสอบการชำระเงิน",
    "rated_by_employer_at": "วันและเวลาที่ผู้ว่าจ้างให้คะแนน",
    "rated_by_worker_at": "วันและเวลาที่ผู้หางานให้คะแนน",
    "rating_by_employer": "คะแนนที่ผู้ว่าจ้างให้ผู้หางาน",
    "rating_by_worker": "คะแนนที่ผู้หางานให้ผู้ว่าจ้าง",
    "reply_to_email": "อีเมลสำหรับรับการตอบกลับ",
    "reply_to_name": "ชื่อผู้รับการตอบกลับ",
    "required_skills_json": "ทักษะจำเป็นของงานในรูปแบบ JSON",
    "resume_file_path": "ตำแหน่งไฟล์ Resume",
    "responded_at": "วันและเวลาที่ตอบรับหรือปฏิเสธคำเชิญ",
    "retired_at": "วันและเวลาที่เลิกใช้ทักษะนี้",
    "review_note": "หมายเหตุจากผู้ตรวจสอบ",
    "reviewed_at": "วันและเวลาที่ตรวจสอบ",
    "reviewed_by_user_id": "รหัสผู้ดูแลที่ตรวจสอบข้อมูล",
    "role": "บทบาทผู้ใช้งาน เช่น ผู้ดูแล ผู้ว่าจ้าง หรือผู้หางาน",
    "sent_at": "วันและเวลาที่ส่งอีเมลสำเร็จ",
    "skill_category_id": "รหัสหมวดทักษะ",
    "skill_id": "รหัสทักษะ",
    "skill_name": "ชื่อทักษะ",
    "sort_order": "ลำดับการแสดงผล",
    "starts_at": "วันและเวลาที่สิทธิ์โปรโมตเริ่มต้น",
    "status": "สถานะการส่งอีเมลในคิว",
    "subject": "หัวข้ออีเมล",
    "submitted_at": "วันและเวลาที่ส่งเอกสาร",
    "token": "ค่าโทเคนแบบสุ่มสำหรับยืนยันตัวตน",
    "token_type": "ประเภทโทเคน เช่น ยืนยันอีเมล หรือตั้งรหัสผ่านใหม่",
    "to_email": "อีเมลผู้รับ",
    "to_name": "ชื่อผู้รับ",
    "updated_at": "วันและเวลาที่แก้ไขข้อมูลล่าสุด",
    "used_at": "วันและเวลาที่ใช้โทเคนแล้ว",
    "user_id": "รหัสผู้ใช้งาน",
    "withdrawn_at": "วันและเวลาที่ผู้หางานถอนใบสมัคร",
    "work_end_date": "วันสิ้นสุดงาน",
    "work_end_time": "เวลาสิ้นสุดงาน",
    "work_interest_id": "รหัสหัวข้องานที่สนใจ",
    "work_location": "สถานที่ทำงานหรือรายละเอียดพื้นที่ปฏิบัติงาน",
    "work_mode": "รูปแบบการทำงาน เช่น หน้างาน ออนไลน์ หรือผสม",
    "work_province": "จังหวัดของสถานที่ทำงานหรือพื้นที่ที่ผู้หางานสะดวก",
    "work_schedule": "ข้อความสรุปวันและเวลาทำงาน",
    "work_start_date": "วันเริ่มงาน",
    "work_start_time": "เวลาเริ่มงาน",
    "worker_profile_id": "รหัสโปรไฟล์ผู้หางาน",
    "worker_user_id": "รหัสผู้ใช้งานของผู้หางาน",
}


def parse_schema(sql: str):
    tables = []
    for match in re.finditer(r"CREATE TABLE `([^`]+)` \((.*?)\) ENGINE=", sql, re.DOTALL):
        name, body = match.groups()
        columns = []
        primary = set()
        unique = set()
        foreign = set()
        for key_match in re.finditer(r"PRIMARY KEY \(([^)]+)\)", body):
            primary.update(re.findall(r"`([^`]+)`", key_match.group(1)))
        for key_match in re.finditer(r"UNIQUE KEY `[^`]+` \(([^)]+)\)", body):
            unique.update(re.findall(r"`([^`]+)`", key_match.group(1)))
        for key_match in re.finditer(r"FOREIGN KEY \(([^)]+)\)", body):
            foreign.update(re.findall(r"`([^`]+)`", key_match.group(1)))
        for line in body.splitlines():
            field_match = re.match(r"\s*`([^`]+)`\s+([^\s,]+)", line)
            if not field_match:
                continue
            field, sql_type = field_match.groups()
            type_name = re.match(r"[a-zA-Z]+", sql_type).group(0).upper()
            size_match = re.search(r"\(([^)]+)\)", sql_type)
            size = size_match.group(1) if size_match else "-"
            if type_name == "ENUM":
                size = "ชุดค่า"
            keys = []
            if field in primary:
                keys.append("PK")
            if field in foreign:
                keys.append("FK")
            if field in unique:
                keys.append("UK")
            columns.append({"field": field, "type": type_name, "size": size, "key": ", ".join(keys) or "-"})
        if columns:
            tables.append((name, columns))
    return tables


def description_for(table: str, field: str) -> str:
    overrides = {
        ("email_log", "id"): "รหัสรายการอีเมลในคิว",
        ("schema_migrations", "migration"): "ชื่อไฟล์ migration ที่นำมาใช้แล้ว",
        ("job_promotions", "amount"): "ยอดเงินที่ต้องชำระสำหรับรายการโปรโมต",
        ("worker_skills", "proficiency_level"): "ระดับความชำนาญของผู้หางานในทักษะนั้น",
        ("employer_documents", "document_status"): "สถานะผลการตรวจเอกสารผู้ว่าจ้าง",
    }
    if (table, field) in overrides:
        return overrides[(table, field)]
    if field in FIELD_DESCRIPTIONS:
        return FIELD_DESCRIPTIONS[field]
    return "ข้อมูล " + field.replace("_", " ")


def set_run_font(run, size, bold=False, color=None):
    """Use the formal Thai report font family used by the reference document."""
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=55, start=80, bottom=55, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    set_run_font(run, size, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_borders(table, color="000000", size="6"):
    """Apply the thin black grid used in the supplied report's data tables."""
    table_pr = table._tbl.tblPr
    borders = table_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, 14)


def build_document():
    tables = parse_schema(SCHEMA.read_text(encoding="utf-8"))
    missing = [name for name, _ in tables if name not in TABLE_INFO]
    if missing:
        raise ValueError("Missing table labels: " + ", ".join(missing))

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "TH Sarabun New"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    normal.font.size = Pt(16)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.name = "TH Sarabun New"
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("3.3.11 Data Dictionary (พจนานุกรมข้อมูล)")
    set_run_font(heading_run, 18, True)

    intro = document.add_paragraph()
    intro.paragraph_format.first_line_indent = Cm(1.25)
    intro.paragraph_format.line_spacing = 1.15
    intro.paragraph_format.space_after = Pt(8)
    intro.add_run(
        "ตารางข้อมูล (Data Table) เป็นการอธิบายรายละเอียดของข้อมูลที่อยู่ในระบบ "
        "โดยระบบได้กำหนดโครงสร้างแฟ้มข้อมูลไว้ในตารางข้อมูล ประกอบด้วย 24 ตาราง "
        "เพื่อใช้จัดเก็บข้อมูลผู้ใช้งาน ผู้ว่าจ้าง ผู้หางาน ประกาศงาน การสมัครงาน และข้อมูลสนับสนุนของระบบ"
    )
    for run in intro.runs:
        set_run_font(run, 16)

    legend = document.add_paragraph()
    legend.paragraph_format.first_line_indent = Cm(1.25)
    legend.paragraph_format.space_after = Pt(14)
    legend_run = legend.add_run("คำย่อ: PK = Primary Key, FK = Foreign Key, UK = Unique Key")
    set_run_font(legend_run, 15)

    widths = [Cm(4.65), Cm(2.40), Cm(1.75), Cm(1.55), Cm(5.20)]
    headers = ["Field Name", "Type", "Size", "Key", "Description"]
    for index, (table_name, columns) in enumerate(tables, start=1):
        english, thai, summary = TABLE_INFO[table_name]
        # Keep dense tables on their own page; short tables flow naturally together,
        # as in the supplied thesis-style reference.
        if index > 1 and len(columns) >= 13:
            document.add_page_break()

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.paragraph_format.space_before = Pt(7)
        caption.paragraph_format.space_after = Pt(4)
        caption.paragraph_format.keep_with_next = True
        caption_run = caption.add_run(f"ตารางที่ 3-{index + 10} รายละเอียดของข้อมูลตาราง {english} ({thai})")
        set_run_font(caption_run, 16, True)

        table = document.add_table(rows=2, cols=5)
        table.autofit = False
        table.style = "Table Grid"
        set_table_borders(table)

        title_cells = table.rows[0].cells
        title_cell = title_cells[0]
        for cell in title_cells[1:]:
            title_cell = title_cell.merge(cell)
        set_cell_text(title_cell, f"{english} ({thai})", 16, False, WD_ALIGN_PARAGRAPH.CENTER)
        prevent_row_split(table.rows[0])

        header = table.rows[1]
        set_repeat_table_header(header)
        for cell, label, width in zip(header.cells, headers, widths):
            cell.width = width
            set_cell_text(cell, label, 15, True, WD_ALIGN_PARAGRAPH.CENTER)

        for row_index, column in enumerate(columns):
            cells = table.add_row().cells
            row = table.rows[-1]
            prevent_row_split(row)
            values = [
                column["field"],
                column["type"],
                column["size"],
                column["key"],
                description_for(table_name, column["field"]),
            ]
            aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
            for cell, value, width, align in zip(cells, values, widths, aligns):
                cell.width = width
                set_cell_text(cell, value, 14, column["key"] == "PK", align)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(header)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Data Dictionary ระบบ FLEXJOB"
    document.core_properties.subject = "รายละเอียดโครงสร้างข้อมูลทุกตาราง"
    document.core_properties.author = "FLEXJOB"
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Tables: {len(tables)}")


if __name__ == "__main__":
    build_document()
